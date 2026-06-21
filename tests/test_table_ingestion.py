from pathlib import Path

from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.ingestion.loader import load_document
from app.ingestion.tables import extract_csv_blocks, extract_xlsx_blocks
from app.ingestion.types import InputType
from app.repositories.documents import DocumentRepository
from tests.helpers import create_authenticated_workspace


def create_sample_csv(path: Path) -> None:
    path.write_text(
        "\n".join(
            [
                "Date,Customer,Amount,Status",
                "2026-01-01,Acme,1200,Paid",
                "2026-01-02,Beta,850,Pending",
                "2026-01-03,Core,990,Paid",
            ]
        ),
        encoding="utf-8",
    )


def create_sample_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sales = workbook.active
    sales.title = "Sales"
    sales.append(["Date", "Customer", "Amount", "Status"])
    sales.append(["2026-01-01", "Acme", 1200, "Paid"])
    sales.append(["2026-01-02", "Beta", 850, "Pending"])
    sales.append(["2026-01-03", "Core", 990, "Paid"])

    inventory = workbook.create_sheet("Inventory")
    inventory.append(["SKU", "Name", "Quantity"])
    inventory.append(["A-1", "Notebook", 12])
    inventory.append(["B-2", "Marker", 30])

    workbook.save(path)


def test_extract_csv_blocks_preserves_headers_and_row_ranges(tmp_path: Path) -> None:
    path = tmp_path / "sample.csv"
    create_sample_csv(path)

    blocks = extract_csv_blocks(path)

    assert len(blocks) == 1

    block = blocks[0]
    assert "CSV Table" in block.text
    assert "Rows: 2-4" in block.text
    assert "Columns: Date | Customer | Amount | Status" in block.text
    assert "Row 2: Date=2026-01-01 | Customer=Acme | Amount=1200 | Status=Paid" in block.text
    assert block.metadata["source_type"] == "csv"
    assert block.metadata["sheet_name"] is None
    assert block.metadata["row_start"] == 2
    assert block.metadata["row_end"] == 4
    assert block.metadata["column_names"] == ["Date", "Customer", "Amount", "Status"]
    assert block.metadata["profile"]["numeric_columns"] == ["Amount"]


def test_extract_csv_blocks_splits_large_tables_into_row_windows(tmp_path: Path) -> None:
    path = tmp_path / "large.csv"
    rows = ["Name,Amount"]
    rows.extend(f"Customer {index},{index * 10}" for index in range(1, 31))
    path.write_text("\n".join(rows), encoding="utf-8")

    blocks = extract_csv_blocks(path)

    assert len(blocks) == 2
    assert blocks[0].metadata["row_start"] == 2
    assert blocks[0].metadata["row_end"] == 26
    assert blocks[1].metadata["row_start"] == 27
    assert blocks[1].metadata["row_end"] == 31
    assert "Rows: 27-31" in blocks[1].text


def test_extract_xlsx_blocks_supports_multiple_sheets(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    create_sample_xlsx(path)

    blocks = extract_xlsx_blocks(path)

    assert {block.metadata["sheet_name"] for block in blocks} == {"Sales", "Inventory"}

    sales = next(block for block in blocks if block.metadata["sheet_name"] == "Sales")
    inventory = next(block for block in blocks if block.metadata["sheet_name"] == "Inventory")

    assert "Sheet: Sales" in sales.text
    assert "Rows: 2-4" in sales.text
    assert sales.metadata["row_start"] == 2
    assert sales.metadata["row_end"] == 4
    assert sales.metadata["column_names"] == ["Date", "Customer", "Amount", "Status"]

    assert "Sheet: Inventory" in inventory.text
    assert "Columns: SKU | Name | Quantity" in inventory.text
    assert inventory.metadata["row_start"] == 2
    assert inventory.metadata["row_end"] == 3
    assert inventory.metadata["profile"]["numeric_columns"] == ["Quantity"]


def test_load_document_routes_csv_and_xlsx(tmp_path: Path) -> None:
    csv_path = tmp_path / "sample.csv"
    xlsx_path = tmp_path / "sample.xlsx"
    create_sample_csv(csv_path)
    create_sample_xlsx(xlsx_path)

    csv_document = load_document(csv_path, InputType.CSV, "sample")
    xlsx_document = load_document(xlsx_path, InputType.XLSX, "sample")

    assert csv_document.source_type == InputType.CSV
    assert csv_document.blocks[0].metadata["column_names"] == [
        "Date",
        "Customer",
        "Amount",
        "Status",
    ]
    assert xlsx_document.source_type == InputType.XLSX
    assert {block.metadata["sheet_name"] for block in xlsx_document.blocks} == {
        "Sales",
        "Inventory",
    }


def test_upload_csv_and_xlsx_succeeds(
    client: TestClient,
    db_session: Session,
    tmp_path: Path,
) -> None:
    workspace_id, headers = create_authenticated_workspace(
        db_session,
        email="table-upload-owner@example.com",
    )
    document_repo = DocumentRepository(db_session)

    csv_path = tmp_path / "upload.csv"
    xlsx_path = tmp_path / "upload.xlsx"
    create_sample_csv(csv_path)
    create_sample_xlsx(xlsx_path)

    for filename, content_type, path in [
        ("upload.csv", "text/csv", csv_path),
        (
            "upload.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx_path,
        ),
    ]:
        response = client.post(
            f"/api/v1/workspaces/{workspace_id}/documents/upload",
            headers=headers,
            files={
                "file": (
                    filename,
                    path.read_bytes(),
                    content_type,
                )
            },
        )

        assert response.status_code == 201
        payload = response.json()
        assert payload["status"] == "succeeded"
        assert payload["chunks_created"] >= 1

        chunks = document_repo.list_chunks_for_document(
            workspace_id=workspace_id,
            document_id=payload["document_id"],
        )
        assert chunks
        assert chunks[0].source_metadata["source_type"] in {"csv", "xlsx"}
        assert "row_start" in chunks[0].source_metadata
